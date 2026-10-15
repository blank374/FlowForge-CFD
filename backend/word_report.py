"""Word evidence report mapped to the supplied CRM-WB standard process."""
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.standard_process import assess


def read(directory, name):
    try:
        return json.loads((directory/name).read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return {}


def write_word_report(directory, result, status='待工程审核'):
    directory = Path(directory)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Normal'].font.size = Pt(10.5)
    document.styles['Normal'].paragraph_format.space_after = Pt(6)
    document.styles['Title'].font.size = Pt(22)

    def paragraph(text):
        document.add_paragraph(str(text))

    def table(headers, rows):
        t = document.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        for cell, value in zip(t.rows[0].cells, headers):
            cell.text = str(value)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EF');cell._tc.get_or_add_tcPr().append(shade)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            for cell, value in zip(t.add_row().cells, values):
                cell.text = str(value if value is not None else '未记录')
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                cell._tc.get_or_add_tcPr().append(borders)

    def heading(text, new_page=False):
        if new_page: document.add_page_break()
        document.add_heading(text, 1)

    def image(name, caption):
        path=directory/name
        if path.is_file():
            p=document.add_paragraph(caption)
            p.paragraph_format.keep_with_next=True
            shape=document.add_picture(str(path), width=Inches(6.4))
            if shape.height > Inches(4.3):
                shape.width=int(shape.width*Inches(4.3)/shape.height)
                shape.height=Inches(4.3)
        else:
            paragraph(caption+'：未生成。')

    geometry=read(directory,'geometry_report.json')
    surface=read(directory,'surface/report.json')
    quality=read(directory,'surface/stl_quality.json')
    units=read(directory,'openfoam_case/units.json') or read(directory,'units.json')
    checkpoint=read(directory,'checkpoint.json')
    standard=assess(directory)
    solver=(directory/'solver.log').read_text(errors='replace') if (directory/'solver.log').is_file() else ''
    courant=[float(x) for x in re.findall(r'Courant Number mean: [0-9.eE+-]+ max: ([0-9.eE+-]+)',solver)]
    clock=re.findall(r'ClockTime = ([0-9.eE+-]+) s',solver)
    mesh_text=''
    for name in ('checkMesh.log','checkMesh_mm.log'):
        if (directory/name).is_file():
            mesh_text=(directory/name).read_text(errors='replace')
            break
    cells_match=re.search(r'(?m)^\s*cells:\s*(\d+)',mesh_text)
    cells=result.get('cells') or (int(cells_match.group(1)) if cells_match else '未提取')
    document.add_heading('流体仿真标准过程与结果报告', 0)
    paragraph('本报告依据《标准过程定义1-WB-20250103结果对比》组织内容，自动记录本次计算的真实参数、网格、数值执行及产物。标准内容覆盖与标准工况验收是不同结论；缺失数据明确标为未完成。')
    table(['项目','本次记录'], [('任务',directory.parent.name),('版本或运行号',directory.name),
          ('输入标识',result.get('input_model','未记录')),('运行状态',status),
          ('CRM 标准工程验收','未完成，不作为通过证明'),('计算范围',result.get('scope','流程测试，尚未验证工程精度'))])
    heading('一 几何造型与处理')
    paragraph('标准指定 DPW4_wb_no_tail_v03.igs；本次输入是否同一外形须按文件及哈希核对。简化翼身或球体不能替代 CRM 标模。')
    table(['字段','记录'], [('工作几何单位',geometry.get('unit','未提取')),('输入 SHA256',geometry.get('source_sha256','未提取')),
          ('Topo 后实体统计',json.dumps(geometry.get('after_topo',{}),ensure_ascii=False)),
          ('Topo 与检查结果',geometry.get('status','未提取')),('人工处理与部件分组','Topo/Paste/Orient 操作细节、wing/body 分组与外形审核记录未完整采集')])
    heading('二 网格定义生成与质量', True)
    table(['项目','本次实测或配置','标准要求'], [
        ('表面目标尺寸',surface.get('target_length_mm','未提取'),'局部线长、点数和分布需记录'),
        ('表面三角形',quality.get('triangles','未提取'),'表面无非预期开放边及相交'),
        ('体单元数',cells,'参考算例约 372 万，不作为固定合格阈值'),
        ('首层高度','未实测','0.01 mm'),('增长率','未实测','原文 <1.2 与 1.2 并存，待确认'),
        ('远场尺寸','未计算半展长比例','>100 半展长，原文绝对长度换算待澄清'),
        ('翼根翼尖及前后缘加密','未验证','<0.1% 半展长或当地弦长；后缘数量原文 >>8 待澄清'),
        ('层数与空间过渡','未实测','1 至 100 层，过渡速率 0.8'),
        ('体积与体积比','见原始网格日志，尚未完整提取','无负体积；体积比一般不超过 1000'),
        ('对称面归零','未验证或当前为全模型','半模型 y=0；边界分组明确'),
        ('单位转换',json.dumps(units,ensure_ascii=False),'明确 mm/m，避免重复缩放')])
    for name in ('checkMesh_mm.log','checkMesh.log'):
        if (directory/name).is_file():
            text=(directory/name).read_text(errors='replace')
            paragraph(name+'：'+('通过普通检查' if 'Mesh OK.' in text else '未通过或执行不完整')+'；警告和详细指标见原日志。')
    heading('三 前处理与计算工况', True)
    paragraph('标准参考工况：M=0.85，Re=5×10⁶，侧滑角 0°，设计 CL=0.5，总温 311 K，总压 201194 Pa，静温 271.73 K，静压 125446.22 Pa；攻角范围的原文表格和正文存在差异，需确认。下表为实际输入，不是标准值的自动替代。')
    physics=result.get('physics',{})
    table(['实际参数','数值'], [(k,v) for k,v in physics.items()] or [('工况','未配置')])
    paragraph('尚需明确马赫数、雷诺数、迎角、参考面积、参考弦长、力矩参考点、湍流闭合与壁面处理，以及入口速度、出口压力和表达式定义。当前 OpenFOAM 字典替代 CFX 的 CCL/DEF；未生成的 CFX 文件不宣称已提供。')
    heading('四 求解监控与收敛')
    table(['字段','记录'], [('求解进程结束','是' if '\nEnd\n' in solver else '否或未提取'),
        ('实际最大 Courant 数',max(courant) if courant else '未提取'),('终止模拟时间 s',result.get('final_time_s','未提取')),
        ('步数',result.get('steps','未提取')),('时间步与配置',json.dumps(checkpoint.get('parameters',{}),ensure_ascii=False)),
        ('数值字典追溯',json.dumps(checkpoint.get('numerics',{}),ensure_ascii=False)),
        ('CL25 与 CL50 收敛',json.dumps(standard['checks'][3]['actual'],ensure_ascii=False) if standard['checks'][3]['actual'] else '未完成，缺少真实 CL/CD/Cm 迭代序列')])
    paragraph('标准要求 |CL25−CL50|<0.0001；还涉及 CD 滑动平均判据，最终采用口径须确认。完成指定时间段、低 Courant 数或程序正常退出均不能代替气动力收敛。当前重试从时间零重新计算，不等于暂停后从已有流场续算。')
    heading('五 后处理结果', True)
    image('p.png','压力云图。若来自不可压求解器，原始 p 通常为运动学压力，单位及参考压力需核对。')
    image('U.png','速度分布。当前切片位置见 ParaView 状态及提取记录。')
    heading('六 残差与实验对比', (directory/'p.png').is_file() or (directory/'U.png').is_file())
    image('residuals.png','残差历史。图像存在不代表残差与气动力均已收敛。')
    paragraph('表面与空间流线：未完成。翼剖面站位、剖面 Cp 曲线及反向纵轴：未完成。Cp 计算须明确静压参考、密度与速度，不能将不可压运动学压力直接当作 Pa。')
    paragraph('DPW 升力、阻力、力矩曲线对比：未完成。设计 CL=0.5 时剖面压力分布对比：未完成。需要相同几何、工况、站位及试验来源，不能将演示模型与 CRM 试验直接作精度验收。Tecplot/Excel 对比数据及误差统计尚未接入自动报告。')
    heading('七 定量考核与过程追溯')
    table(['字段','记录'], [('表面脚本耗时 s',surface.get('elapsed_seconds','未提取')),
        ('求解日志末次墙钟耗时 s',clock[-1] if clock else '未提取'),('操作人员熟练度与人工耗时','待填写'),
        ('CPU GPU 内存与并行核数','待采集，原标准硬件数据不作为本机数据'),('导入导出耗时与并行效率','待采集'),
        ('核时与费用','未计算，需实际资源与计费依据')])
    paragraph('原始证据：输入文件、geometry_report.json、surface/report.json、Pointwise 原生工程、checkMesh 日志、solver.log、checkpoint.json、ParaView 状态和结果字段。失败和重试保留各自版本，报告不替代原始记录。')
    heading('八 标准要求覆盖与结论')
    table(['项目','证据状态'], [(c['label'],{'MISSING':'缺少证据','FAILED':'不符合或未通过','PASSED':'该项证据通过'}[c['status']]) for c in standard['checks']])
    paragraph('结论：本报告已按标准过程列示内容，但当前证据不足以认定 CRM-WB 标准工况或工程精度通过。未完成项应补齐后重新出具；审核人、日期与审批结论待填写。')
    destination=directory/'report.docx'
    temporary=directory/'report.tmp.docx'
    for root in (document.element, document.styles.element):
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    document.save(temporary)
    temporary.replace(destination)
    return destination
