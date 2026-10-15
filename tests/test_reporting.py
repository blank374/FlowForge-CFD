import base64
import pytest
from backend.reporting import write_report

def test_report_requires_real_images_and_escapes_input(tmp_path):
    data={'input_model':'<script>bad</script>','cells':100,'physics':{'solver':'OpenFOAM 13'}}
    with pytest.raises(RuntimeError):write_report(tmp_path,data)
    for name in ['p.png','U.png','residuals.png']:(tmp_path/name).write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a8nUAAAAASUVORK5CYII='))
    write_report(tmp_path,data)
    content=(tmp_path/'report.html').read_text(encoding='utf-8')
    assert '<script>bad' not in content and '&lt;script&gt;' in content
    assert content.count('data:image/png;base64,')==3
    from docx import Document
    document=Document(tmp_path/'report.docx')
    text='\n'.join(p.text for p in document.paragraphs)
    assert '标准工况验收' in text
    assert 'DPW' in text and '未完成' in text
    assert len(document.inline_shapes)==3
    assert all('<script>bad' not in p.text for p in document.paragraphs if 'input_model' in p.text)


def test_word_report_missing_evidence_does_not_invent_results(tmp_path):
    from backend.word_report import write_word_report
    from docx import Document
    path=write_word_report(tmp_path,{'input_model':'未运行模型','physics':{}},status='失败记录')
    document=Document(path)
    text='\n'.join(p.text for p in document.paragraphs)
    assert '不足以认定' in text and '表面与空间流线：未完成' in text
    assert len(document.inline_shapes)==0
